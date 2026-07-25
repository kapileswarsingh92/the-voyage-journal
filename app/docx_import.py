"""Parses an uploaded Word (.docx) document into the same real-HTML shape
the rich story editor already stores, plus any embedded images — used by
the "Insert file" toolbar button's smarter behavior for .docx uploads
(see app/blog.py's docx_import() route and static/js/main.js's
handleDocxImport()).

Only the modern, XML-based .docx format is supported here (via
python-docx) — legacy binary .doc files aren't parseable this way and
still go through the plain "attach as a download card" path, same as PDF
and Pages files. This is a read-only, best-effort conversion: python-docx
exposes paragraphs/runs/images cleanly, but tables, headers/footers,
footnotes, and most advanced Word formatting have no equivalent in this
site's intentionally small rich-text vocabulary (bold, italic, lists,
font size on headings) and are simply dropped rather than approximated.

The returned HTML is later run through the same sanitize-on-submit
pipeline as any other story content (see utils.sanitize_story_html /
blog.normalize_story_content) once the user actually submits the story —
so a bug here can produce an odd-looking import, but can't itself become
an XSS vector.
"""

import html as html_lib

from docx import Document
from docx.oxml.ns import qn

from .utils import MAX_PDF_BYTES


class DocxImportError(ValueError):
    """Raised with a message safe to show the user directly."""


# Word's built-in heading styles get a bit of visual weight on import
# (bold + one of the site's existing font-size tiers) since the editor
# has no dedicated heading tag — anything not listed here still gets
# bolded via the generic "heading*" prefix check below, just without an
# enlarged size.
_HEADING_SIZES = {"heading 1": "28px", "title": "28px", "heading 2": "22px"}


def _run_html(run) -> str:
    """A single run's text, escaped and wrapped in <b>/<i> per its bold/
    italic flags. Word can also set bold/italic at the paragraph-style
    level rather than per-run; that's intentionally not chased here — it's
    a rare case for the kind of prose this feature targets, and getting it
    wrong silently would be worse than a run occasionally importing
    unstyled."""
    text = html_lib.escape(run.text or "")
    if not text:
        return ""
    if run.bold:
        text = f"<b>{text}</b>"
    if run.italic:
        text = f"<i>{text}</i>"
    return text


def _paragraph_image_parts(doc, para):
    """Yields each embedded image's document Part (in run order) for every
    <w:drawing>/<a:blip> found in this paragraph's runs."""
    for run in para.runs:
        for blip in run._element.findall(".//" + qn("a:blip")):
            rId = blip.get(qn("r:embed"))
            if not rId:
                continue
            try:
                yield doc.part.related_parts[rId]
            except KeyError:
                continue  # broken/external reference — skip rather than fail the whole import


def parse_docx(file_storage, max_images: int):
    """Parse an uploaded .docx into (html, images).

    html contains the story body as real HTML in this site's existing
    vocabulary (<p>, <b>, <i>, <ul><li>, and a <span style="font-size">
    wrapper for imported headings), with a [[docximg:N]] placeholder
    paragraph (1-indexed) wherever an embedded image appeared in the
    document — the caller (the /docx-import route) never sees a real
    [[photo:N]] token here, since these images haven't been through the
    normal upload/save pipeline yet; the client turns each placeholder
    into a real inserted photo block itself.

    images is a list of {"content_type": str, "data": bytes} in the same
    order the placeholders reference.

    Raises DocxImportError (message is safe to show the user) if the file
    can't be read, is empty, or contains more images than max_images.
    """
    file_storage.stream.seek(0, 2)
    size = file_storage.stream.tell()
    file_storage.stream.seek(0)
    if size > MAX_PDF_BYTES:
        raise DocxImportError(
            f'"{file_storage.filename}" is larger than the {MAX_PDF_BYTES // (1024 * 1024)}MB limit.'
        )

    try:
        doc = Document(file_storage.stream)
    except Exception:
        raise DocxImportError(
            "Couldn't read that Word document — it may be corrupted, password-protected, "
            "or not actually a .docx file."
        )

    html_parts = []
    images = []
    list_buffer = []

    def flush_list():
        if list_buffer:
            html_parts.append("<ul>" + "".join(f"<li>{item}</li>" for item in list_buffer) + "</ul>")
            list_buffer.clear()

    for para in doc.paragraphs:
        style_name = (para.style.name or "").strip().lower() if para.style else ""
        is_list_item = style_name.startswith("list")
        inline_html = "".join(_run_html(r) for r in para.runs)

        new_images = []
        for part in _paragraph_image_parts(doc, para):
            if len(images) + len(new_images) >= max_images:
                raise DocxImportError(
                    f"This document has more than {max_images} images, which is this story's photo "
                    "limit — remove some images from the document and try again."
                )
            new_images.append({"content_type": part.content_type, "data": part.blob})

        # A list-style paragraph that also happens to carry an image is
        # rare enough to just fall through to normal paragraph handling
        # (which flushes any open list first) rather than special-casing it.
        if is_list_item and not new_images:
            if inline_html:
                list_buffer.append(inline_html)
            continue
        flush_list()

        if inline_html:
            size_px = _HEADING_SIZES.get(style_name)
            if size_px:
                html_parts.append(f'<p><span style="font-size:{size_px}"><b>{inline_html}</b></span></p>')
            elif style_name.startswith("heading") or style_name == "title":
                html_parts.append(f"<p><b>{inline_html}</b></p>")
            else:
                html_parts.append(f"<p>{inline_html}</p>")

        for img in new_images:
            images.append(img)
            html_parts.append(f"<p>[[docximg:{len(images)}]]</p>")

    flush_list()

    if not images and not "".join(html_parts).strip():
        raise DocxImportError("That Word document doesn't seem to have any readable text or images.")

    return "".join(html_parts), images
